from flask import Flask, request, jsonify, render_template
from werkzeug.utils import secure_filename
import os
from PIL import Image
import pytesseract
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

import pdfplumber
from docx import Document
from google.cloud import dialogflow_v2 as dialogflow
from google.cloud.dialogflow_v2.types import TextInput, QueryInput
import uuid
import json
import re
from fpdf import FPDF
from google.protobuf.json_format import MessageToDict
import ast
import string
from flask import send_from_directory
from flask import send_file
from langdetect import detect, LangDetectException

app = Flask(__name__)
UPLOAD_FOLDER = 'uploads'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

os.environ["GOOGLE_APPLICATION_CREDENTIALS"] = "ma_clef_dialogflow.json"
PROJECT_ID = "monprojetdialogflow-465522"

def extract_text(filepath):
    text = ""
    if filepath.endswith('.pdf'):
        with pdfplumber.open(filepath) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    text += t + "\n"
    elif filepath.endswith('.docx'):
        doc = Document(filepath)
        text = "\n".join([para.text for para in doc.paragraphs])
    elif filepath.lower().endswith(('.png', '.jpg', '.jpeg')):
        img = Image.open(filepath)
        text = pytesseract.image_to_string(img)
    else:
        raise ValueError("Unsupported file type")
    
    try:
        lang = detect(text)
    except LangDetectException:
        lang = "unknown"

    return text.strip(),lang

def clean_feature(feat):
    # Enlever ponctuation finale (ex: '.'), et mettre en minuscules
    return feat.strip().rstrip(string.punctuation).lower()

feature_mapping = {
    "formulaire de contact": "formulaire de contact",
    "contact form": "formulaire de contact",
    "contact": "formulaire de contact",

    "map": "carte interactive",
    "interactive map": "carte interactive",
    "carte google": "carte interactive",
    "google maps": "carte interactive",
    "carte interactive": "carte interactive",

    "gallery": "portfolio",
    "photo gallery": "portfolio",
    "galerie": "portfolio",
    "galerie photo": "portfolio",
    "portfolio": "portfolio",

    "recruitment": "recrutement",
    "recrutement": "recrutement",
    "job page": "recrutement",
    "careers": "recrutement",

    "blog": "blog",
    "articles": "blog",

    "home": "page d'accueil",
    "homepage": "page d'accueil",
    "page d'accueil": "page d'accueil",
    "accueil": "page d'accueil",

    "about": "page à propos",
    "about us": "page à propos",
    "à propos": "page à propos",
    "page à propos": "page à propos",

    "services": "page services",
    "our services": "page services",
    "page services": "page services",
    "service": "page services",

    "products": "page produits",
    "page produits": "page produits",

    "faq": "faq",
    "frequently asked questions": "faq",
    "foire aux questions": "faq",

    "testimonials": "témoignages",
    "reviews": "témoignages",
    "témoignages": "témoignages",
    "clients": "témoignages",

    "legal": "pages légales",
    "legal pages": "pages légales",
    "mentions légales": "pages légales",

    "social media": "réseaux sociaux",
    "facebook": "réseaux sociaux",
    "instagram": "réseaux sociaux",
    "youtube": "réseaux sociaux",
    "réseaux sociaux": "réseaux sociaux",

    "responsive design": "responsive design",
    "responsive": "responsive design",

    "security": "sécurité et performance",
    "performance": "sécurité et performance",
    "security and performance": "sécurité et performance",
    "sécurité": "sécurité et performance",
    "performance": "sécurité et performance",
}

def clean_features(features_raw):
    articles = {
        # Français
        'un', 'une', 'le', 'la', 'les', 'des', 'nos', 'mes', 'notre', 'votre', 'leur', 'son', 'sa', 'ses',

        # Anglais
        'a', 'an', 'the', 'my', 'your', 'our', 'their', 'his', 'her', 'its'
    }
    cleaned = []
    seen = set()

    for f in features_raw:
        key = f.lower().strip()
        words = key.split()
        if words and words[0] in articles:
            words = words[1:]
        key = " ".join(words)
        key = feature_mapping.get(key, key)  
        if key not in seen:
            cleaned.append(key)
            seen.add(key)
    return cleaned

def convert_words_to_digits(text):
    mapping = {
    # Français
    'un': '1', 'une': '1', 'deux': '2', 'trois': '3', 'quatre': '4',
    'cinq': '5', 'six': '6', 'sept': '7', 'huit': '8', 'neuf': '9',
    'dix': '10', 'dizaine de': '10', 'onze': '11', 'douze': '12', 'treize': '13',
    'quatorze': '14', 'quinze': '15', 'seize': '16',
    'dix-sept': '17', 'dix-huit': '18', 'dix-neuf': '19', 'vingt': '20',

    # Anglais
    'one': '1', 'two': '2', 'three': '3', 'four': '4', 'five': '5',
    'six': '6', 'seven': '7', 'eight': '8', 'nine': '9', 'ten': '10',
    'eleven': '11', 'twelve': '12','dozen' : '12', 'thirteen': '13', 'fourteen': '14',
    'fifteen': '15', 'sixteen': '16', 'seventeen': '17', 'eighteen': '18',
    'nineteen': '19', 'twenty': '20'
}


    pattern = r'\b(' + '|'.join(mapping.keys()) + r')\b'
    return re.sub(pattern, lambda m: mapping[m.group(0).lower()], text, flags=re.IGNORECASE)

def extract_number_before_pages(text):
    match = re.search(r'(\d+)\s+(pages)', text.lower())
    if match:
        return int(match.group(1))
    return 0

def detect_intent_text(text):
    session_client = dialogflow.SessionsClient()
    session_id = str(uuid.uuid4())
    session = session_client.session_path(PROJECT_ID, session_id)
    text_input = TextInput(text=text, language_code='fr')
    query_input = QueryInput(text=text_input)
    response = session_client.detect_intent(
        request={"session": session, "query_input": query_input}
    )

    result = response.query_result
    fulfillment = result.fulfillment_text
    params = result.parameters
    
    website_type = params.get("website_type", "Non spécifié")
    try:
        number_of_pages = int(params.get("number_of_pages", 0))
    except Exception:
        number_of_pages = 0
    
    if number_of_pages == 0:
       number_of_pages = extract_number_before_pages(text)
       print(f"[Fallback regex] Nombre de pages détecté : {number_of_pages}")


    raw_features = params.get("website_feature")
    print("RAW FEATURES FROM DIALOGFLOW:", raw_features)

    

    # Nettoyage : enlever ponctuation finale et mettre en minuscules
    features = [clean_feature(f) for f in raw_features]

    

    # --- Enrichir les features manquants en analysant le texte brut ---
    all_possible_features = list(feature_mapping.keys())
    text_lower = text.lower()

    for feature in all_possible_features:
        if feature in text_lower and feature not in features:
            features.append(feature)

    features = clean_features(features)
    print("Features list:", features)
    print("Nombre d'éléments dans features:", len(features))

    if number_of_pages == 0:
        number_of_pages = len(features)
        print("Nombre de pages estimé par le nombre de fonctionnalités:", number_of_pages)


    return fulfillment, number_of_pages, features, website_type
# --------------- Partie de prix -------------


with open('prices.json', 'r', encoding='utf-8') as f:
    prices = json.load(f)

# Calculate total price :
def calculate_price(number_of_pages, features):
    total = 0

    for tranche in prices['pages']:
        if tranche['min'] <= number_of_pages <= tranche['max']:
            total += tranche['price']
            break

    for feat in features:
        key = feature_mapping.get(feat) 
        if key and key in prices['features']:
            total += prices['features'][key]

    return total

# Generating the pdf:
def generate_quote_pdf(client_text, number_of_pages, features, total_price, website_type, filename='devis.pdf', site_language='fr'):
    pdf = FPDF()
    pdf.add_page()

    # Couleur de fond sombre
    pdf.set_fill_color(45, 45, 45)
    pdf.rect(0, 0, pdf.w, pdf.h, 'F')
    pdf.set_draw_color(255, 255, 255)
    pdf.set_text_color(255, 255, 255)

    pdf.set_font("Helvetica", 'B', 30)
    pdf.set_xy(10, 15)

    title = "DEVIS" if site_language == 'fr' else "QUOTE"
    pdf.cell(0, 10, title, ln=True)

    pdf.image('static/images/logo.png', x=150, y=13, w=40)
    pdf.ln(25)

    pdf.set_font("Arial", '', 12)
    if site_language == 'fr':
        pdf.cell(0, 10, "Commencez à voir des résultats - plus tôt nous commençons, plus vite vous gagnez..", ln=True)
    else:
        pdf.cell(0, 10, "Start seeing results - the sooner we start, the faster you grow.", ln=True)
    pdf.ln(10)

    # Table headers
    col1_width, col2_width, col3_width = 60, 80, 40
    line_height = 10

    pdf.set_font("Arial", 'B', 12)
    pdf.set_text_color(245, 180, 0)
    if site_language == 'fr':
        headers = ["Élément", "Détail", "Prix"]
    else:
        headers = ["Item", "Detail", "Price"]
    pdf.cell(col1_width, line_height, headers[0], border=1, fill=True, align='C')
    pdf.cell(col2_width, line_height, headers[1], border=1, fill=True, align='C')
    pdf.cell(col3_width, line_height, headers[2], border=1, ln=True, fill=True, align='C')

    # Données
    pdf.set_font("Arial", '', 12)
    pdf.set_text_color(255, 255, 255)

    if site_language == 'fr':
        pdf.cell(col1_width, line_height, "Type de site", border=1, align='C')
    else:
        pdf.cell(col1_width, line_height, "Website type", border=1, align='C')
    pdf.cell(col2_width, line_height, website_type, border=1, align='C')
    pdf.cell(col3_width, line_height, "-", border=1, ln=True, align='C')

    for tranche in prices['pages']:
        if tranche['min'] <= number_of_pages <= tranche['max']:
            page_price = tranche['price']
            break
    else:
        page_price = "Inconnu" if site_language == 'fr' else "Unknown"

    if site_language == 'fr':
        pdf.cell(col1_width, line_height, "Nombre de pages", border=1, align='C')
        pdf.cell(col2_width, line_height, f"{number_of_pages} page(s)", border=1, align='C')
    else:
        pdf.cell(col1_width, line_height, "Number of pages", border=1, align='C')
        pdf.cell(col2_width, line_height, f"{number_of_pages} page(s)", border=1, align='C')
    pdf.cell(col3_width, line_height, f"{page_price} dh", border=1, ln=True, align='C')

    # Fonctionnalités
    if features:
        for feat in features:
            key = feature_mapping.get(feat.lower())
            prix = prices['features'].get(key, "Inconnu") if key else "Inconnu"
            feat_label = "Fonctionnalité" if site_language == 'fr' else "Feature"
            pdf.cell(col1_width, line_height, feat_label, border=1, align='C')
            pdf.cell(col2_width, line_height, feat, border=1, align='C')
            pdf.cell(col3_width, line_height, f"{prix} dh", border=1, ln=True, align='C')
    else:
        feat_label = "Fonctionnalité" if site_language == 'fr' else "Feature"
        no_feat = "Aucune" if site_language == 'fr' else "None"
        pdf.cell(col1_width, line_height, feat_label, border=1, align='C')
        pdf.cell(col2_width, line_height, no_feat, border=1, align='C')
        pdf.cell(col3_width, line_height, "-", border=1, ln=True, align='C')

    # Total
    pdf.set_font("Arial", 'B', 13)
    pdf.cell(col1_width + col2_width, line_height, "Total", border=1, align='C')
    pdf.cell(col3_width, line_height, f"{total_price} dh", border=1, ln=True, align='C')

    # Footer
    pdf.ln(5)
    pdf.set_font("Arial", '', 10)
    if site_language == 'fr':
        pdf.cell(182, 10, "Confirmez ce devis et nous nous occupons du reste.", ln=True, align='R')
    else:
        pdf.cell(182, 10, "Confirm this quote and we'll take care of the rest.", ln=True, align='R')

    output_path = os.path.join("uploads", filename)
    pdf.output(output_path)
    return output_path


@app.route('/')
def index():
    return render_template('index.html') 

@app.route('/fr')
def index_fr():
    return render_template('index_fr.html')

@app.route("/demande_de_devis")
def demande_de_devis():
    return render_template("demande_de_devis.html")

@app.route("/demande_de_devis_fr")
def demande_de_devis_fr():
    return render_template("demande_de_devis_fr.html") 

def clean_text(text):
    import re
    text = re.sub(r'\n+', ' ', text)   # remplace retours à la ligne multiples par un espace
    text = re.sub(r'\s+', ' ', text)   # remplace espaces multiples par un seul espace
    return text.strip() 
def generate_confirmation_message(site_type, number_of_pages, features , site_language):
        
    feature_translation = {
            "formulaire de contact": "contact form",
            "recrutement": "recruitment",
            "carte interactive": "interactive map",
            "blog": "blog",
            "portfolio": "portfolio",
            "page d'accueil": "homepage",
            "page à propos": "about page",
            "page services": "services page",
            "page produits": "products page",
            "faq": "faq",
            "témoignages": "testimonials",
            "pages légales": "legal pages",
            "réseaux sociaux": "social media",
            "responsive design": "responsive design",
            "sécurité et performance": "security and performance"
        }

    if site_language == 'en':
        features = [feature_translation.get(f.lower(), f) for f in features]
        
    feat_str = ", ".join(features)

    if site_language == 'fr':
        return f"Bien compris, vous souhaitez créer un {site_type} de {number_of_pages} pages avec les fonctionnalités suivantes : {feat_str}. Confirmez-vous ce devis ?"
    else :
        return f"Got it! You want to create a {site_type} with {number_of_pages} pages and the following features: {feat_str}. Do you confirm this quote?"


@app.route('/upload', methods=['POST'])
def upload():
    try:
        file = request.files.get('file')
        if not file:
            return jsonify({"error": "Aucun fichier reçu"}), 400
        
        site_language = request.form.get("site_language", "en")

        filename = secure_filename(file.filename)
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(filepath)

        text , detected_lang = extract_text(filepath)

        # Blocage si langue différente
        if detected_lang != site_language:
            if site_language == 'fr':
                return jsonify({
                    "error": f"Le document est en '{detected_lang}', mais le site est en '{site_language}'. Veuillez télécharger un fichier dans la même langue."
                }), 400
            else:
                return jsonify({
                    "error": f"The document language is in '{detected_lang}', but the site is in '{site_language}'. Please upload a file in the same language."
                }), 400


        text = clean_text(text)
        text = convert_words_to_digits(text) 
        reply, number_of_pages, features, website_type = detect_intent_text(text)
        
        features = clean_features(features)

        print("Nombre de pages extrait:", number_of_pages)

        total_price = calculate_price(number_of_pages, features)


        return jsonify({
            "CLIENT": text,
            "CHENOCOM": generate_confirmation_message(website_type, number_of_pages, features, site_language),
            "TYPE_SITE": website_type,
            "number_of_pages": number_of_pages,
            "FEATURES": features,
            "WAITING_CONFIRMATION": True ,
            "detected_language": detected_lang,
            "site_language": site_language
        })
        """ return jsonify({
            "pdf_filename": "devis.pdf",
            "message": "Your quote is ready !"
    }) """
    
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    
@app.route('/confirm', methods=['POST'])
def confirm():
    try:
        data = request.json
        confirmation = data.get("confirmation")
        site_language = data.get("site_language", "en")

        if confirmation.lower() != "oui":
            message = "Veuillez reformuler votre demande." if site_language == 'fr' else "Please reformulate your request."
            return jsonify({"message": message})

        client_text = data["CLIENT"]
        number_of_pages = data["number_of_pages"]
        features = data["FEATURES"]
        website_type = data["TYPE_SITE"]

        if site_language == "en":
            feature_translation = {
                "formulaire de contact": "contact form",
                "recrutement": "recruitment",
                "carte interactive": "interactive map",
                "blog": "blog",
                "portfolio": "portfolio",
                "page d'accueil": "homepage",
                "page à propos": "about page",
                "page services": "services page",
                "page produits": "products page",
                "faq": "faq",
                "témoignages": "testimonials",
                "pages légales": "legal pages",
                "réseaux sociaux": "social media",
                "responsive design": "responsive design",
                "sécurité et performance": "security and performance"
            }
            features = [feature_translation.get(f.lower(), f) for f in features]
                

        total_price = calculate_price(number_of_pages, features)
        pdf_path = generate_quote_pdf(client_text, number_of_pages, features, total_price, website_type, site_language=site_language)

        message = "Devis généré avec succès." if site_language == 'fr' else "Quote generated successfully."

        return jsonify({
            "message": message,
            "download_url": f"/download?filename={os.path.basename(pdf_path)}"
        })

    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/download')
def download_file():
    filename = request.args.get('filename')
    if not filename:
        return "Nom de fichier manquant", 400
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename, as_attachment=True)

if __name__ == '__main__':
    print("+++++++++++ Flask is running at http://127.0.0.1:5000/ ++++++")
    app.run(debug=True)




#Activate venv : venv\Scripts\activate
# Run : python app.py
